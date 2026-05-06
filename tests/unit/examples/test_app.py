from __future__ import annotations

import io

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("docx")

from fastapi.testclient import TestClient

from examples.redaction_studio import app as app_module
from examples.redaction_studio.app import app


def _make_docx(paragraphs: list[list[str]]) -> bytes:
    import docx
    from docx.enum.text import WD_BREAK

    document = docx.Document()
    for page_idx, lines in enumerate(paragraphs):
        for line_idx, line in enumerate(lines):
            para = document.add_paragraph(line)
            if page_idx < len(paragraphs) - 1 and line_idx == len(lines) - 1:
                para.add_run().add_break(WD_BREAK.PAGE)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


@pytest.fixture(autouse=True)
def _clear_store():
    app_module.STORE._docs.clear()
    yield
    app_module.STORE._docs.clear()


@pytest.fixture
def client():
    return TestClient(app)


def test_upload_returns_doc_id_and_page_count(client):
    raw = _make_docx([["alpha"], ["beta"]])
    resp = client.post(
        "/api/upload",
        files={"file": ("sample.docx", raw, "application/octet-stream")},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["filename"] == "sample.docx"
    assert body["pageCount"] == 2
    assert isinstance(body["docId"], str) and len(body["docId"]) >= 8


def test_upload_rejects_unsupported_extension(client):
    resp = client.post(
        "/api/upload",
        files={"file": ("nope.txt", b"hi", "text/plain")},
    )
    assert resp.status_code == 400
    assert "Unsupported" in resp.json()["detail"]


def test_list_pages_returns_text(client):
    raw = _make_docx([["page one text"], ["page two text"]])
    doc_id = client.post(
        "/api/upload",
        files={"file": ("a.docx", raw, "application/octet-stream")},
    ).json()["docId"]

    resp = client.get(f"/api/documents/{doc_id}")
    assert resp.status_code == 200
    body = resp.json()
    assert body["pageCount"] == 2
    assert body["pages"][0]["text"].startswith("page one")
    assert body["pages"][1]["index"] == 1


def test_list_pages_404_when_unknown(client):
    resp = client.get("/api/documents/missing")
    assert resp.status_code == 404
