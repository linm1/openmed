from __future__ import annotations

import io

from .types import RedactedPage, UploadedDoc

try:
    import docx as _docx
    from docx.enum.text import WD_BREAK as _WD_BREAK

    DOCX_AVAILABLE = True
except ImportError:
    _docx = None
    _WD_BREAK = None
    DOCX_AVAILABLE = False

try:
    from reportlab.pdfgen.canvas import Canvas as _Canvas
    from reportlab.lib.pagesizes import LETTER as _LETTER

    PDF_AVAILABLE = True
except ImportError:
    _Canvas = None
    _LETTER = None
    PDF_AVAILABLE = False


def _final_text_for(page_index: int, doc: UploadedDoc, redacted: dict[int, RedactedPage]) -> str:
    if page_index in redacted:
        return redacted[page_index].redacted
    return doc.pages[page_index].text


def write_docx(doc: UploadedDoc, redacted: dict[int, RedactedPage]) -> bytes:
    if not DOCX_AVAILABLE:
        raise RuntimeError("python-docx not installed.")
    out = _docx.Document()
    for i, _page in enumerate(doc.pages):
        body = _final_text_for(i, doc, redacted)
        for line in body.split("\n"):
            out.add_paragraph(line)
        if i < len(doc.pages) - 1:
            out.add_paragraph().add_run().add_break(_WD_BREAK.PAGE)
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()


def write_pdf(doc: UploadedDoc, redacted: dict[int, RedactedPage]) -> bytes:
    raise NotImplementedError  # filled in next task
